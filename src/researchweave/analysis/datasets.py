"""Immutable imported datasets, durable revision store and reviewed variables."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
from collections.abc import Callable
from dataclasses import replace
from pathlib import Path
from typing import Any, Literal, cast
from zipfile import BadZipFile

from .errors import AnalysisError, require
from .models import Dataset, Variable, canonical_json, digest, utc_now

MAX_BYTES = 50 * 1024 * 1024
MAX_ROWS = 500_000
MAX_COLUMNS = 2000
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_DECIMAL_COMMA = re.compile(r"^[+-]?\d+,\d+(?:[eE][+-]?\d+)?$")


def validate_dataset_identity(dataset_id: str, revision: str | None = None) -> None:
    """Identifiers are hex digests, never filesystem paths or user-supplied labels."""
    require(
        isinstance(dataset_id, str)
        and len(dataset_id) == 24
        and all(c in "0123456789abcdef" for c in dataset_id),
        "Invalid dataset ID; expected 24 lowercase hex characters",
    )
    if revision is not None:
        require(
            isinstance(revision, str)
            and len(revision) == 64
            and all(c in "0123456789abcdef" for c in revision),
            "Invalid revision; expected 64 lowercase hex characters",
        )


def expected_revision(dataset: Dataset) -> str:
    """Reconstruct content-derived dataset revisions from supported constructors."""
    record_list = list(dataset.records)
    variables = [v.__dict__ for v in dataset.variables]
    operation = dataset.operation
    if dataset.parent_revision is None:
        if operation.get("type") == "import_csv":
            payload = {
                "id": dataset.dataset_id,
                "source": dataset.source_hash,
                "rows": record_list,
                "variables": variables,
                "operation": operation,
            }
        elif operation.get("type") == "import_table":
            payload = {
                "id": dataset.dataset_id,
                "source": dataset.source_hash,
                "rows": record_list,
                "variables": variables,
                "operation": operation,
            }
        else:
            raise AnalysisError("INTEGRITY_FAILURE", "Unknown root revision constructor")
    elif operation.get("type") == "join":
        payload = {
            "columns": dataset.columns,
            "rows": record_list,
            "variables": variables,
            "operation": operation,
        }
    elif operation.get("type") == "survey_indicator":
        payload = {"operation": operation, "records": record_list}
    else:
        payload = {
            "parent": dataset.parent_revision,
            "records": record_list,
            "variables": variables,
            "operation": operation,
        }
    return digest(payload)


def import_csv(
    path: Path,
    *,
    encoding: str = "utf-8-sig",
    delimiter: str | None = None,
    decimal: str = ".",
    grouping: str | None = None,
    missing: tuple[str, ...] = ("",),
    dataset_id: str | None = None,
) -> Dataset:
    require(decimal in (".", ","), "decimal must be dot or comma")
    require(
        grouping is None or (len(grouping) == 1 and grouping != decimal),
        "Invalid grouping separator",
    )
    require(path.is_file(), "Import source does not exist", "NOT_FOUND")
    require(
        path.stat().st_size <= MAX_BYTES, "Source exceeds configured byte budget", "RESOURCE_LIMIT"
    )
    content = path.read_bytes()
    if encoding == "auto":
        if content.startswith((b"\xff\xfe", b"\xfe\xff")):
            encoding = "utf-16"
        elif content.startswith(b"\xef\xbb\xbf"):
            encoding = "utf-8-sig"
        else:
            encoding = "utf-8"
    try:
        text = content.decode(encoding)
    except (UnicodeError, LookupError) as exc:
        raise AnalysisError(
            "INVALID_INPUT", "Text decoding failed; select another encoding"
        ) from exc
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:16_384], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    require(len(delimiter) == 1, "Delimiter must contain exactly one character")
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter, restkey="__overflow__")
    require(bool(reader.fieldnames), "CSV must include header")
    columns = tuple(reader.fieldnames or ())
    require(
        len(columns) == len(set(columns)) and all(columns), "Duplicate/empty columns prohibited"
    )
    require("__overflow__" not in columns, "Reserved column name used")
    require(len(columns) <= MAX_COLUMNS, "Too many columns", "RESOURCE_LIMIT")
    rows: list[dict[str, Any]] = []
    for row in reader:
        require(len(rows) < MAX_ROWS, "Too many rows", "RESOURCE_LIMIT")
        require(
            "__overflow__" not in row and all(k in columns for k in row),
            "CSV row has extra columns",
            "SCHEMA_MISMATCH",
        )
        require(
            all(value is not None for value in row.values()),
            "CSV row has missing fields",
            "SCHEMA_MISMATCH",
        )
        normalized: dict[str, Any] = {}
        for column in columns:
            raw = row[column]
            if raw in missing:
                normalized[column] = None
            else:
                candidate = raw.strip()
                if grouping is not None:
                    candidate = candidate.replace(grouping, "")
                if decimal == "," and _DECIMAL_COMMA.fullmatch(candidate):
                    candidate = candidate.replace(",", ".")
                normalized[column] = candidate if candidate != raw.strip() else raw
        rows.append(normalized)
    variables = tuple(Variable(name=c, missing_values=missing) for c in columns)
    source_hash = hashlib.sha256(content).hexdigest()
    identifier = dataset_id or digest({"source_hash": source_hash})[:24]
    validate_dataset_identity(identifier)
    op = {
        "type": "import_csv",
        "encoding": encoding,
        "delimiter": delimiter,
        "decimal": decimal,
        "grouping": grouping,
        "missing": list(missing),
    }
    revision = digest(
        {
            "id": identifier,
            "source": source_hash,
            "rows": rows,
            "variables": [v.__dict__ for v in variables],
            "operation": op,
        }
    )
    return Dataset(
        identifier, revision, None, source_hash, utc_now(), columns, variables, tuple(rows), op
    )


def export_csv(
    dataset: Dataset,
    destination: Path,
    *,
    delimiter: str = ",",
    decimal: str = ".",
    formula_safe: bool = True,
) -> dict[str, Any]:
    """Write a deterministic CSV snapshot without silently overwriting existing files."""

    require(len(delimiter) == 1, "Delimiter must be one character")
    require(decimal in {".", ","}, "Decimal must be dot or comma")
    require(destination.parent.is_dir(), "Export directory missing", "NOT_FOUND")
    require(not destination.exists(), "Export destination exists", "CONFLICT")
    require(destination.suffix.lower() in {".csv", ".tsv"}, "CSV export suffix required")

    def cell(value: Any) -> str:
        if value is None:
            return ""
        text = str(value)
        if decimal == "," and isinstance(value, (int, float)):
            text = text.replace(".", ",")
        if formula_safe and text.startswith(_FORMULA_PREFIXES):
            return "'" + text
        return text

    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(
        buffer,
        fieldnames=list(dataset.columns),
        delimiter=delimiter,
        lineterminator="\n",
        quoting=csv.QUOTE_MINIMAL,
    )
    writer.writeheader()
    for row in dataset.records:
        writer.writerow({column: cell(row[column]) for column in dataset.columns})
    payload = buffer.getvalue().encode("utf-8-sig")
    require(len(payload) <= MAX_BYTES * 2, "Export exceeds byte budget", "RESOURCE_LIMIT")
    destination.write_bytes(payload)
    return {
        "path": str(destination),
        "sha256": hashlib.sha256(payload).hexdigest(),
        "bytes": len(payload),
        "delimiter": delimiter,
        "decimal": decimal,
        "encoding": "utf-8-sig",
        "formula_safe": formula_safe,
        "warning": (
            "Spreadsheet-formula safety prefixes formula-like text with an apostrophe, "
            "which intentionally changes those exported cell strings"
            if formula_safe
            else "Formula-like cells are preserved exactly and may execute in spreadsheet software"
        ),
    }


def import_table(
    path: Path,
    *,
    format: str | None = None,
    sheet: int | str = 0,
    dataset_id: str | None = None,
    data_only: bool = True,
    min_row: int = 1,
    max_row: int | None = None,
    min_col: int = 1,
    max_col: int | None = None,
) -> Dataset:
    """Import nonexecuting structured formats through opt-in vetted readers."""
    import pandas as pd

    require(path.is_file(), "Import source does not exist", "NOT_FOUND")
    require(path.stat().st_size <= MAX_BYTES, "Source exceeds byte budget", "RESOURCE_LIMIT")
    kind = (format or path.suffix.lower().lstrip(".")).lower()
    require(min_row >= 1 and min_col >= 1, "Workbook ranges use positive one-based indices")
    require(max_row is None or max_row >= min_row, "Invalid workbook row range")
    require(max_col is None or max_col >= min_col, "Invalid workbook column range")
    if kind in ("csv", "tsv"):
        return import_csv(path, delimiter="\t" if kind == "tsv" else None, dataset_id=dataset_id)
    if kind == "parquet":
        try:
            import pyarrow.parquet as pq
        except ImportError as exc:
            raise AnalysisError(
                "DEPENDENCY_MISSING", "Reader for parquet is not installed"
            ) from exc
        try:
            parquet = pq.ParquetFile(path)
            metadata = parquet.metadata
            require(
                metadata.num_rows <= MAX_ROWS and metadata.num_columns <= MAX_COLUMNS,
                "Parquet metadata exceeds row/column budget",
                "RESOURCE_LIMIT",
            )
            names = tuple(str(name) for name in parquet.schema_arrow.names)
            require(
                len(names) == len(set(names)) and all(names),
                "Duplicate/empty columns prohibited",
            )
            rows: list[dict[str, Any]] = []
            for batch in parquet.iter_batches(batch_size=65_536):
                batch_rows = batch.to_pylist()
                require(
                    len(rows) + len(batch_rows) <= MAX_ROWS,
                    "Parquet row budget exceeded during scan",
                    "RESOURCE_LIMIT",
                )
                rows.extend({str(key): value for key, value in row.items()} for row in batch_rows)
            numeric_prefixes = (
                "int",
                "uint",
                "float",
                "double",
                "halffloat",
                "decimal",
            )
            labels = tuple(
                Variable(
                    name=name,
                    dtype=(
                        "number"
                        if str(parquet.schema_arrow.field(index).type).startswith(numeric_prefixes)
                        else "text"
                    ),
                )
                for index, name in enumerate(names)
            )
            source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
            identifier = dataset_id or digest({"source": source_hash})[:24]
            validate_dataset_identity(identifier)
            operation = {
                "type": "import_table",
                "format": kind,
                "reader": "pyarrow.iter_batches",
                "batch_size": 65_536,
            }
            revision = digest(
                {
                    "id": identifier,
                    "source": source_hash,
                    "rows": rows,
                    "variables": [v.__dict__ for v in labels],
                    "operation": operation,
                }
            )
            return Dataset(
                identifier,
                revision,
                None,
                source_hash,
                utc_now(),
                names,
                labels,
                tuple(rows),
                operation,
            )
        except AnalysisError:
            raise
        except (ValueError, TypeError, OSError) as exc:
            raise AnalysisError("INVALID_INPUT", "Could not parse parquet source") from exc
    elif kind in {"sav", "zsav", "por", "dta", "xpt"}:
        try:
            import pyreadstat
        except ImportError as exc:
            raise AnalysisError(
                "DEPENDENCY_MISSING", f"Reader for {kind} is not installed"
            ) from exc
        readstat = cast(Callable[..., tuple[Any, Any]], {
            "sav": pyreadstat.read_sav,
            "zsav": pyreadstat.read_sav,
            "por": pyreadstat.read_por,
            "dta": pyreadstat.read_dta,
            "xpt": pyreadstat.read_xport,
        }[kind])
        readstat_options: dict[str, Any] = (
            {} if kind == "xpt" else {"apply_value_formats": False}
        )
        if kind in {"sav", "zsav"}:
            # Preserve SPSS user-missing values in the data so the dictionary can
            # distinguish them from system missing values instead of silently
            # converting them to NaN during import.
            readstat_options["user_missing"] = True
        try:
            frame, metadata = readstat(path, **readstat_options)
        except (ValueError, TypeError, OSError) as exc:
            raise AnalysisError("INVALID_INPUT", f"Could not parse {kind} source") from exc
        labels_by_name = getattr(metadata, "column_names_to_labels", {}) or {}
        value_labels = getattr(metadata, "variable_value_labels", {}) or {}
        missing_ranges = getattr(metadata, "missing_ranges", {}) or {}
        missing_user_values = getattr(metadata, "missing_user_values", {}) or {}
        measure = getattr(metadata, "variable_measure", {}) or {}
        metadata_variables = {
            str(name): {
                "label": str(labels_by_name.get(name) or ""),
                "value_labels": {
                    (
                        f"{float(key):.17g}"
                        if isinstance(key, (int, float)) and not isinstance(key, bool)
                        else str(key)
                    ): str(value)
                    for key, value in (value_labels.get(name, {}) or {}).items()
                },
                "missing_values": tuple(
                    [str(value) for value in (missing_user_values.get(name, []) or [])]
                    + [
                        (
                            str(entry.get("lo"))
                            if isinstance(entry, dict) and entry.get("lo") == entry.get("hi")
                            else f"{entry.get('lo')}..{entry.get('hi')}"
                            if isinstance(entry, dict)
                            and "lo" in entry
                            and "hi" in entry
                            else str(entry)
                        )
                        for entry in (missing_ranges.get(name, []) or [])
                    ]
                ),
                "scale": (
                    "ordinal"
                    if str(measure.get(name) or "unknown").lower() == "ordinal"
                    else "nominal"
                    if str(measure.get(name) or "unknown").lower() == "nominal"
                    else "interval"
                    if str(measure.get(name) or "unknown").lower() == "scale"
                    else "unknown"
                ),
            }
            for name in frame.columns
        }
    else:
        metadata_variables = None
        readers = {
            "ods": lambda: pd.read_excel(path, sheet_name=sheet, engine="odf"),
            "json": lambda: pd.read_json(path),
            "jsonl": lambda: pd.read_json(path, lines=True),
        }
        if kind == "xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise AnalysisError(
                    "DEPENDENCY_MISSING", "Reader for xlsx is not installed"
                ) from exc
            try:
                workbook = load_workbook(path, read_only=True, data_only=data_only)
                if isinstance(sheet, int):
                    require(
                        0 <= sheet < len(workbook.sheetnames),
                        "Workbook sheet index out of range",
                    )
                    worksheet = workbook[workbook.sheetnames[sheet]]
                else:
                    require(
                        sheet in workbook.sheetnames, "Workbook sheet does not exist", "NOT_FOUND"
                    )
                    worksheet = workbook[str(sheet)]
                selected_max_row = min(max_row or worksheet.max_row, worksheet.max_row)
                selected_max_col = min(max_col or worksheet.max_column, worksheet.max_column)
                require(
                    selected_max_row - min_row + 1 <= MAX_ROWS + 1
                    and selected_max_col - min_col + 1 <= MAX_COLUMNS,
                    "Workbook range exceeds row/column budget",
                    "RESOURCE_LIMIT",
                )
                values = list(
                    worksheet.iter_rows(
                        min_row=min_row,
                        max_row=selected_max_row,
                        min_col=min_col,
                        max_col=selected_max_col,
                        values_only=True,
                    )
                )
                require(bool(values), "Workbook range is empty")
                header = ["" if value is None else str(value) for value in values[0]]
                require(
                    all(header) and len(header) == len(set(header)),
                    "Workbook header has duplicates/empties",
                )
                frame = pd.DataFrame(values[1:], columns=header)
            except AnalysisError:
                raise
            except (ValueError, TypeError, OSError, KeyError, BadZipFile) as exc:
                raise AnalysisError("INVALID_INPUT", "Could not parse xlsx source") from exc
        else:
            require(kind in readers, f"Unsupported import format: {kind}", "UNSUPPORTED_FORMAT")
            try:
                frame = readers[kind]()
                if kind == "ods":
                    require(
                        isinstance(frame, pd.DataFrame),
                        "Workbook sheet selection did not resolve to one table",
                    )
                    selected_max_row = min(max_row or (len(frame) + 1), len(frame) + 1)
                    selected_max_col = min(max_col or len(frame.columns), len(frame.columns))
                    require(
                        selected_max_row >= min_row and selected_max_col >= min_col,
                        "Workbook range is empty",
                    )
                    require(
                        selected_max_row - min_row + 1 <= MAX_ROWS + 1
                        and selected_max_col - min_col + 1 <= MAX_COLUMNS,
                        "Workbook range exceeds row/column budget",
                        "RESOURCE_LIMIT",
                    )
                    # pandas already consumed row 1 as the header. Interpret
                    # workbook row/column bounds with the same one-based semantics
                    # as XLSX: row 1 is the header and data starts at row 2.
                    start_data = max(min_row - 2, 0)
                    stop_data = max(selected_max_row - 1, 0)
                    start_col = min_col - 1
                    stop_col = selected_max_col
                    frame = frame.iloc[start_data:stop_data, start_col:stop_col].copy()
            except ImportError as exc:
                raise AnalysisError(
                    "DEPENDENCY_MISSING", f"Reader for {kind} is not installed"
                ) from exc
            except (ValueError, TypeError, OSError) as exc:
                raise AnalysisError("INVALID_INPUT", f"Could not parse {kind} source") from exc
    require(
        len(frame) <= MAX_ROWS and len(frame.columns) <= MAX_COLUMNS,
        "Parsed table exceeds row/column budget",
        "RESOURCE_LIMIT",
    )
    names = tuple(str(c) for c in frame.columns)
    require(len(set(names)) == len(names), "Duplicate columns prohibited")
    rows = json.loads(frame.to_json(orient="records", date_format="iso"))
    imported_variables: list[Variable] = []
    for i, c in enumerate(names):
        rich = (metadata_variables or {}).get(c, {})
        rich_label = str(rich.get("label", ""))
        rich_value_labels = cast(dict[str, str], rich.get("value_labels", {}))
        rich_missing = cast(tuple[str, ...], rich.get("missing_values", ()))
        rich_scale = cast(
            Literal["nominal", "ordinal", "interval", "ratio", "unknown"],
            rich.get("scale", "unknown"),
        )
        imported_variables.append(
            Variable(
                name=c,
                dtype="number" if pd.api.types.is_numeric_dtype(frame.iloc[:, i]) else "text",
                label=rich_label,
                value_labels=rich_value_labels,
                missing_values=rich_missing,
                scale=rich_scale,
            )
        )
    labels = tuple(imported_variables)
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    identifier = dataset_id or digest({"source": source_hash})[:24]
    validate_dataset_identity(identifier)
    operation = {
        "type": "import_table",
        "format": kind,
        "sheet": sheet,
        "data_only": data_only if kind == "xlsx" else None,
        "range": (
            {"min_row": min_row, "max_row": max_row, "min_col": min_col, "max_col": max_col}
            if kind in {"xlsx", "ods"}
            else None
        ),
        "metadata_limitations": (
            ["POR metadata exposed by pyreadstat may omit column labels/value labels"]
            if kind == "por"
            else []
        ),
    }
    revision = digest(
        {
            "id": identifier,
            "source": source_hash,
            "rows": rows,
            "variables": [v.__dict__ for v in labels],
            "operation": operation,
        }
    )
    return Dataset(
        identifier, revision, None, source_hash, utc_now(), names, labels, tuple(rows), operation
    )


def scan_parquet_page(
    path: Path,
    *,
    offset: int = 0,
    limit: int = 100,
    columns: list[str] | None = None,
    batch_size: int = 65_536,
) -> dict[str, Any]:
    """Read one bounded Parquet page without materializing the complete table.

    The scan validates file-level row/column budgets from Parquet metadata, then
    converts only the requested slice of each Arrow batch to Python objects.
    This path is intended for large-table inspection before an analyst chooses
    whether to create an in-memory immutable Dataset revision.
    """

    require(path.is_file(), "Parquet source does not exist", "NOT_FOUND")
    require(path.stat().st_size <= MAX_BYTES, "Source exceeds byte budget", "RESOURCE_LIMIT")
    require(offset >= 0 and 1 <= limit <= 10_000, "Invalid Parquet page range")
    require(1 <= batch_size <= 65_536, "Parquet batch size outside supported range")
    try:
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise AnalysisError("DEPENDENCY_MISSING", "Reader for parquet is not installed") from exc
    try:
        parquet = pq.ParquetFile(path)
        metadata = parquet.metadata
        require(
            metadata.num_rows <= MAX_ROWS and metadata.num_columns <= MAX_COLUMNS,
            "Parquet metadata exceeds row/column budget",
            "RESOURCE_LIMIT",
        )
        available = [str(name) for name in parquet.schema_arrow.names]
        require(
            len(available) == len(set(available)) and all(available),
            "Duplicate/empty columns prohibited",
        )
        selected = available if columns is None else [str(value) for value in columns]
        require(bool(selected), "At least one Parquet column required")
        require(len(selected) == len(set(selected)), "Duplicate Parquet scan columns")
        require(set(selected) <= set(available), "Unknown Parquet scan column", "NOT_FOUND")
        require(len(selected) <= MAX_COLUMNS, "Too many selected columns", "RESOURCE_LIMIT")
        rows: list[dict[str, Any]] = []
        seen = 0
        converted_rows = 0
        for batch in parquet.iter_batches(batch_size=batch_size, columns=selected):
            batch_rows = batch.num_rows
            batch_start = seen
            batch_end = seen + batch_rows
            seen = batch_end
            if batch_end <= offset:
                continue
            local_start = max(offset - batch_start, 0)
            remaining = limit - len(rows)
            if remaining <= 0:
                break
            local_length = min(batch_rows - local_start, remaining)
            if local_length <= 0:
                continue
            sliced = batch.slice(local_start, local_length)
            converted = sliced.to_pylist()
            converted_rows += len(converted)
            rows.extend(
                {str(key): value for key, value in record.items()} for record in converted
            )
            if len(rows) >= limit:
                break
        schema = {
            name: str(parquet.schema_arrow.field(available.index(name)).type) for name in selected
        }
        return {
            "format": "parquet",
            "reader": "pyarrow.iter_batches_slice",
            "offset": offset,
            "limit": limit,
            "total_rows": int(metadata.num_rows),
            "total_columns": int(metadata.num_columns),
            "columns": selected,
            "schema": schema,
            "rows": rows,
            "returned_rows": len(rows),
            "converted_python_rows": converted_rows,
            "batch_size": batch_size,
            "materialized_entire_table": False,
        }
    except AnalysisError:
        raise
    except (ValueError, TypeError, OSError) as exc:
        raise AnalysisError("INVALID_INPUT", "Could not scan parquet source") from exc


def revise(
    dataset: Dataset,
    records: list[dict[str, Any]],
    operation: dict[str, Any],
    variables: tuple[Variable, ...] | None = None,
) -> Dataset:
    metadata = variables or dataset.variables
    require(len(records) <= MAX_ROWS, "Too many transformed rows", "RESOURCE_LIMIT")
    require(
        all(set(row) == set(dataset.columns) for row in records),
        "Transformed schema mismatch",
        "SCHEMA_MISMATCH",
    )
    identity = digest(
        {
            "parent": dataset.revision,
            "records": records,
            "variables": [v.__dict__ for v in metadata],
            "operation": operation,
        }
    )
    return Dataset(
        dataset.dataset_id,
        identity,
        dataset.revision,
        dataset.source_hash,
        utc_now(),
        dataset.columns,
        metadata,
        tuple(records),
        operation,
    )


def revise_schema(
    dataset: Dataset,
    records: list[dict[str, Any]],
    columns: tuple[str, ...],
    variables: tuple[Variable, ...],
    operation: dict[str, Any],
) -> Dataset:
    """Create an immutable revision whose schema intentionally changes."""

    require(len(records) <= MAX_ROWS, "Too many transformed rows", "RESOURCE_LIMIT")
    require(bool(columns) and len(set(columns)) == len(columns), "Invalid transformed columns")
    require(tuple(variable.name for variable in variables) == columns, "Variable schema mismatch")
    require(
        all(set(row) == set(columns) for row in records),
        "Transformed schema mismatch",
        "SCHEMA_MISMATCH",
    )
    identity = digest(
        {
            "parent": dataset.revision,
            "records": records,
            "variables": [v.__dict__ for v in variables],
            "operation": operation,
        }
    )
    return Dataset(
        dataset.dataset_id,
        identity,
        dataset.revision,
        dataset.source_hash,
        utc_now(),
        columns,
        variables,
        tuple(records),
        operation,
    )


def set_variable(dataset: Dataset, name: str, **updates: Any) -> Dataset:
    require(name in dataset.columns, "Unknown variable", "NOT_FOUND")
    allowed = {
        "label",
        "units",
        "scale",
        "value_labels",
        "missing_values",
        "category_order",
        "allowed_min",
        "allowed_max",
        "privacy_class",
        "private",
        "dtype",
    }
    require(set(updates) <= allowed, "Unsupported metadata property")
    if "scale" in updates:
        require(
            updates["scale"] in {"nominal", "ordinal", "interval", "ratio", "unknown"},
            "Invalid measurement scale",
        )
    if "privacy_class" in updates:
        require(
            updates["privacy_class"] in {"public", "internal", "sensitive", "restricted"},
            "Invalid privacy class",
        )
    if "category_order" in updates:
        category_order = tuple(str(value) for value in updates["category_order"])
        require(len(category_order) == len(set(category_order)), "Duplicate category order entry")
        updates["category_order"] = category_order
    if "missing_values" in updates:
        updates["missing_values"] = tuple(str(value) for value in updates["missing_values"])
    if "allowed_min" in updates or "allowed_max" in updates:
        current = next(variable for variable in dataset.variables if variable.name == name)
        lower = updates.get("allowed_min", current.allowed_min)
        upper = updates.get("allowed_max", current.allowed_max)
        require(
            lower is None or upper is None or float(lower) <= float(upper),
            "Variable allowed range is inverted",
        )
    metadata = tuple(
        replace(variable, **updates) if variable.name == name else variable
        for variable in dataset.variables
    )
    return revise(
        dataset,
        list(dataset.records),
        {"type": "variable_edit", "name": name, "changes": updates},
        metadata,
    )


def page(dataset: Dataset, *, offset: int = 0, limit: int = 100) -> dict[str, Any]:
    require(offset >= 0 and 1 <= limit <= 1000, "Invalid page range")
    return {
        "total": len(dataset.records),
        "offset": offset,
        "limit": limit,
        "rows": list(dataset.records[offset : offset + limit]),
        "revision": dataset.revision,
    }


def data_view(
    dataset: Dataset,
    *,
    offset: int = 0,
    limit: int = 100,
    sort_by: str | None = None,
    descending: bool = False,
    search: str | None = None,
) -> dict[str, Any]:
    """Stable row identities with deterministic filtering, sorting and paging."""

    require(offset >= 0 and 1 <= limit <= 1000, "Invalid page range")
    require(sort_by is None or sort_by in dataset.columns, "Unknown sort column")
    require(search is None or len(search) <= 1000, "Search text exceeds budget", "RESOURCE_LIMIT")
    indexed = list(enumerate(dataset.records))
    if search:
        needle = search.casefold()
        indexed = [
            (index, row)
            for index, row in indexed
            if any(value is not None and needle in str(value).casefold() for value in row.values())
        ]
    if sort_by is not None:
        indexed.sort(
            key=lambda item: (
                item[1][sort_by] is None,
                "" if item[1][sort_by] is None else str(item[1][sort_by]).casefold(),
                item[0],
            ),
            reverse=descending,
        )
    selected = indexed[offset : offset + limit]
    return {
        "total": len(indexed),
        "source_total": len(dataset.records),
        "offset": offset,
        "limit": limit,
        "sort_by": sort_by,
        "descending": descending,
        "search": search,
        "rows": [
            {
                "row_id": digest(
                    {
                        "dataset_id": dataset.dataset_id,
                        "revision": dataset.revision,
                        "source_index": index,
                    }
                )[:24],
                "source_index": index,
                "values": dict(row),
            }
            for index, row in selected
        ],
        "revision": dataset.revision,
    }


class DatasetStore:
    def __init__(self, directory: Path):
        self.root = directory.resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def save(self, dataset: Dataset) -> Path:
        validate_dataset_identity(dataset.dataset_id, dataset.revision)
        require(
            expected_revision(dataset) == dataset.revision,
            "Dataset revision checksum does not match content",
            "INTEGRITY_FAILURE",
        )
        directory = self.root / dataset.dataset_id
        directory.mkdir(parents=True, exist_ok=True)
        target = directory / f"{dataset.revision}.json"
        record = dataset.as_json()
        content = canonical_json(
            {"schema_version": 2, "dataset": record, "content_sha256": digest(record)}
        )
        if target.exists():
            require(
                target.read_bytes() == content, "Existing revision differs", "INTEGRITY_FAILURE"
            )
            return target
        temporary = directory / f".{dataset.revision}.{os.getpid()}.tmp"
        try:
            with temporary.open("xb") as handle:
                handle.write(content)
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)
        return target

    def load(self, dataset_id: str, revision: str) -> Dataset:
        validate_dataset_identity(dataset_id, revision)
        path = self.root / dataset_id / f"{revision}.json"
        require(path.is_file(), "Revision not found", "NOT_FOUND")
        require(
            path.stat().st_size <= MAX_BYTES * 4,
            "Stored dataset exceeds read budget",
            "RESOURCE_LIMIT",
        )
        try:
            envelope = json.loads(path.read_bytes())
        except (ValueError, UnicodeError) as exc:
            raise AnalysisError("INTEGRITY_FAILURE", "Saved revision is not valid JSON") from exc
        require(
            envelope.get("schema_version") == 2 and isinstance(envelope.get("dataset"), dict),
            "Unsupported or unauthenticated dataset snapshot",
            "INTEGRITY_FAILURE",
        )
        obj = envelope["dataset"]
        require(
            digest(obj) == envelope.get("content_sha256"),
            "Saved dataset checksum mismatch",
            "INTEGRITY_FAILURE",
        )
        require(
            obj["revision"] == revision and obj["dataset_id"] == dataset_id,
            "Revision identity mismatch",
            "INTEGRITY_FAILURE",
        )
        result = Dataset(
            dataset_id,
            revision,
            obj["parent_revision"],
            obj["source_hash"],
            obj["created_at"],
            tuple(obj["columns"]),
            tuple(
                Variable(
                    **{
                        **v,
                        "missing_values": tuple(v.get("missing_values", ())),
                        "category_order": tuple(v.get("category_order", ())),
                    }
                )
                for v in obj["variables"]
            ),
            tuple(obj["records"]),
            obj["operation"],
        )
        require(
            expected_revision(result) == revision,
            "Revision does not match restored dataset",
            "INTEGRITY_FAILURE",
        )
        return result
